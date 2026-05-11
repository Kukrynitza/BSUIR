from io import BytesIO


class DocxParser:
    def parse(self, content: bytes) -> str:
        try:
            from docx import Document
        except ImportError as e:
            raise RuntimeError(
            ) from e

        doc = Document(BytesIO(content))
        parts: list[str] = []
        for para in doc.paragraphs:
            t = (para.text or "").strip()
            if t:
                parts.append(t)
        for table in doc.tables:
            for row in table.rows:
                cells = [(c.text or "").strip() for c in row.cells]
                row_text = " ".join(x for x in cells if x)
                if row_text:
                    parts.append(row_text)
        return "\n".join(parts) if parts else ""
